"""Gate for the model-laid-out review deck renderer (local/deck2pptx.mjs
v1.2, over lib/deckspec.mjs + lib/designsystem.mjs + draft2pptx v1.3's
exported emitter + svg2pptx v1.5's parseFigureSvg).

Renders a deck spec over a representative TestPlanGen draft and reads
the .pptx back with python-pptx, asserting the deck contract:

  1. the package opens; the slide walk is the spec's order minus the
     dropped slide; the dropped slide is named on stderr with its
     finding; the machine banner reaches no slide
  2. design system on the slide: title runs at Fluent title1 × 1.5
     (36 pt), body2 items at 18 pt, caption labels at 13.5 pt; Segoe UI
     on every run; title / section / closing on ink, the rest on paper;
     the footer (plan title + page number) on paper slides only
  3. native objects: the steps as checkbox shapes + text, the Expected
     Result / Trace as filled cards, the Coverage Map as a native table
     with its cells, the flow as homePlate + chevron autoshapes with
     the step text INSIDE the shape, the stats tiles with the
     deterministic counts, the comparison panels with their pills
  4. figures: a cited story SVG under --media lands as ONE native shape
     group (the svg2pptx emitter) inside the figure region; a
     generated SVG under --figures likewise; an inline figurespec is
     grounded and rendered; a figure that resolves nowhere degrades to
     a muted "(not embedded)" note and stderr names the fix
  5. speaker notes: a slide's "notes" become a native notes page
  6. [VERIFY: …] spans surface as amber runs inside copied text
  7. CLI contract: default sibling name (<stem>--deck.pptx), -o, --spec
     with a sentinel-wrapped reply file, usage errors exit nonzero, a
     --media / --figures path that is not a directory is refused, a
     spec with no surviving slide exits 1
  8. --generate: ONE streamed model call to a mock Anthropic endpoint
     with the repo prompt verbatim ({PlanTitle}/{Draft}/{Figures}
     substituted, testplangen.deckMaxTokens as max_tokens), the spec
     written beside the deck as <out>.deck.json and re-renderable with
     --spec; a sentinel-less reply exits nonzero and writes nothing
  9. --design (v1.1): the same spec on carbon renders in IBM Plex Sans
     with Carbon's heading-05 title (36 pt, regular weight), the Gray
     100 inverse and square cards; on uswds in Public Sans with the
     size-13 title (40.5 pt bold) and the primary-darker inverse; the
     provenance line names the design; an unknown design is refused
     up front; --generate takes the design from testplangen.deckDesign
 10. --theme (v1.2): carbon dark puts paper slides on Gray 100 with
     Gray 10 text, the native table's header on the divider surface,
     and the embedded figure group re-coloured (no Diagram Style
     Framework ink left in it; Carbon's Blue 50 in its place) — on
     fluent light the group is byte-identical to the unstyled one;
     an unknown theme is refused; the provenance names the theme

Needs python-pptx (review/harness/requirements.txt — the CI
full-format job installs it). Usage: python3 check_deck2pptx.py
"""
import json
import os
import re
import subprocess
import sys
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

HERE = os.path.dirname(os.path.abspath(__file__))
JOB = os.path.join(HERE, "..", "deck2pptx.mjs")
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(f"  {'ok  ' if cond else 'FAIL'} {name}" + ("" if cond else f"  <- {str(detail)[:400]}"))


FIG_URL = "https://mock.example/sites/lrsworkspace/LRS%20Doc%20Index/media/doc12_slide2_fig1.svg"
FIG_ALT = "Routes R1 and R2 before the merge"

DRAFT = f"""<!-- machine-generated test-plan draft — TestPlanGen prompt v1.12 · local/testplangen.mjs v1.16 · provider anthropic -->
> [!WARNING]
> **DRAFT — machine-generated, unreviewed.** Generated 2026-09-05T00:00:00Z from user story doc 12 — "Route Merge". Source sidecar: <https://example/sc.md>
> Review every case and resolve all [VERIFY] items before use.

<!-- verify: 1 finding(s) -->
> [!IMPORTANT]
> Draft verifier: 1 finding(s) — review these first:
> - TC-N1 carries a **Trace:** line

# Test Plan — Route Merge

## Overview

| Surface | Target release | PE |
| --- | --- | --- |
| Pro | 3.8 | Claire Wang |

Verifies measure-preserving merge of two routes in ArcGIS Pro.

## Setup / Prerequisites
- [ ] 1. LRS network with two mergeable routes. [VERIFY: minimum network configuration]
- [x] 2. Two Pro sessions signed in.

## Positive Tests

### TC-P1 — Merge preserves measures
**Steps:**
- [ ] 1. Run Merge Routes on route R1 and route R2.
- [ ] 2. Inspect the measures on the merged route R1.

**Expected Result:** The merged route keeps the source measures unchanged.

**Trace:** "the merge must preserve measures" — story requirement.

**Figure:** ![{FIG_ALT}]({FIG_URL})

## Negative Tests

> [!CAUTION]
> A pass below is the described denial or error — never the edit succeeding.

### TC-N1 — Merge denied on locked route
**Steps:**
- [ ] 1. As User A, lock route R1.
- [ ] 2. As User B, attempt Merge Routes on route R1.

**Expected Result:** The merge is denied with a lock conflict.

**Trace:** exemplar pattern — multi-user denial case.

## Open Questions
- [ ] [VERIFY: minimum network configuration for setup]

## Coverage Map

| # | Requirement (source) | Covered by |
| --- | --- | --- |
| 1 | "the merge must preserve measures" (requirement) | TC-P1 |
| 2 | denial on locked routes (conflict statement) | TC-N1 |

## Generated Figures

_Deterministic addendum — figures PROPOSED by the TestPlanFigures prompt v0.1._

### TC-N1 — Merge denied on locked route

![User A locks R1; User B's merge is denied](<draft--fig-tc-n1.svg>)
"""

# the SlideFigures-vocabulary story figure (plate + routes + ticks + texts)
FIG_SVG = (
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 760 320" width="760" height="320" role="img" aria-label="Slide 2 diagram">'
    '<title>Slide 2 — Routes before the merge</title><desc>Routes R1 and R2 end to end.</desc>'
    "<style>.plate{fill:#FFFFFF;stroke:#D7DFDF;stroke-width:1}.ln{fill:none;stroke-linecap:round;stroke-linejoin:round}"
    ".route{stroke:#16302F;stroke-width:3;stroke-linecap:butt}.tick{stroke:#6E8285;stroke-width:1.15}"
    ".measure{font-size:11px;fill:#6E8285}.id{font-size:12.5px;font-weight:600}.f-cool{fill:#1B6E8C}"
    "text{font-family:'Segoe UI',Arial,sans-serif}</style>"
    '<rect class="plate" x="1" y="1" width="758" height="318" rx="6"/><g transform="translate(20,40)">'
    '<line class="ln route" x1="40" y1="60" x2="340" y2="60"/><line class="ln route" x1="360" y1="60" x2="660" y2="60"/>'
    '<line class="ln tick" x1="40" y1="52" x2="40" y2="68"/><text class="measure" x="40" y="44" text-anchor="middle">0</text>'
    '<text class="id f-cool" x="190" y="84" text-anchor="middle">R1</text><text class="id f-cool" x="510" y="84" text-anchor="middle">R2</text>'
    "</g></svg>")

SPEC = {
    "plan": "Test Plan — Route Merge",
    "slides": [
        {"pattern": "title", "regions": {"headline": "Route Merge", "facts": [{"label": "Surface", "value": "Pro"}, {"label": "PE", "value": "Claire Wang"}]},
         "notes": "Open with the story context; the PE owns the plan."},
        {"pattern": "stats", "title": "At a glance", "eyebrow": "Overview", "regions": {
            "tiles": [{"count": "positive-cases", "label": "Positive cases", "tone": "success"},
                      {"count": "negative-cases", "label": "Negative cases", "tone": "danger"},
                      {"count": "verify-flags", "label": "Open [VERIFY] flags", "tone": "warning"},
                      {"count": "coverage-rows", "label": "Requirements traced", "tone": "brand"}],
            "lede": {"from": {"overview": "prose"}},
            "callout": {"label": "Draft verifier", "body": {"from": {"overview": "verify"}}, "tone": "warning"}}},
        {"pattern": "checklist", "title": "Setup / Prerequisites", "regions": {"items": {"from": {"section": "Setup / Prerequisites"}}}},
        {"pattern": "section", "tone": "success", "regions": {"number": "01", "headline": "Positive Tests", "strap": "one case"}},
        {"pattern": "two-column", "title": "Merge preserves measures", "eyebrow": "TC-P1 · Positive", "tone": "success", "regions": {
            "left": {"from": {"case": "TC-P1", "field": "steps"}},
            "right": [{"label": "Expected result", "body": {"from": {"case": "TC-P1", "field": "expected"}}, "tone": "success"},
                      {"label": "Trace", "body": {"from": {"case": "TC-P1", "field": "trace"}}, "tone": "brand"}]}},
        {"pattern": "figure", "title": "Routes before the merge", "source": "TC-P1", "regions": {"figure": "doc12_slide2_fig1.svg", "aside": ["Run Merge Routes on route R1 and route R2."]}},
        {"pattern": "flow", "title": "Lock, then merge", "source": "TC-N1", "tone": "danger", "regions": {
            "steps": ["As User A, lock route R1.", "As User B, attempt Merge Routes on route R1.", "The merge is denied with a lock conflict."],
            "outcome": {"label": "Expected", "body": {"from": {"case": "TC-N1", "field": "expected"}}, "tone": "danger"}}},
        {"pattern": "figure", "title": "The lock, as a sequence", "source": "TC-N1", "regions": {"figure": "draft--fig-tc-n1.svg"}},
        {"pattern": "comparison", "title": "Unlocked vs locked", "regions": {
            "left": {"label": "TC-P1", "tone": "success", "items": ["The merged route keeps the source measures unchanged."]},
            "right": {"label": "TC-N1", "tone": "danger", "items": ["The merge is denied with a lock conflict."]}}},
        {"pattern": "cards", "title": "Invented", "regions": {"cards": [{"label": "A", "body": "This never appears in the draft."}, {"label": "B", "body": "One route remains."}]}},
        {"pattern": "table", "title": "Coverage Map", "tone": "brand", "regions": {"table": {"from": {"section": "Coverage Map"}}}},
        {"pattern": "figure", "title": "Inline spec", "source": "TC-P1", "regions": {"figure": {"spec": {
            "case": "TC-P1", "rule": "R3", "kind": "topology", "title": "TC-P1 — Merge preserves measures", "caption": "R1 and R2 merge.",
            "nodes": [{"id": "R1", "label": "R1", "shape": "box", "tone": "cool"}, {"id": "R2", "label": "R2", "shape": "box", "tone": "warm"}],
            "edges": [{"from": "R1", "to": "R2", "label": "merge", "style": "solid", "arrow": True}],
            "source": {"steps": [1], "expected": False, "tables": []}}}}},
        {"pattern": "closing", "regions": {"headline": "Before this draft becomes the plan", "asks": ["Review both cases against the story", "Resolve the open [VERIFY] flag"]}},
    ],
}

DECK_REPLY = "Here is the deck.\n[[[DECK BEGIN]]]\n" + json.dumps({"plan": "Test Plan — Route Merge", "slides": SPEC["slides"][:5] + [SPEC["slides"][-1]]}) + "\n[[[DECK END]]]\n"


class MockState:
    def __init__(self):
        self.calls = 0
        self.last_body = {}
        self.text = DECK_REPLY


def make_handler(state):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_POST(self):
            n = int(self.headers.get("content-length") or 0)
            body = json.loads(self.rfile.read(n) if n else b"{}")
            state.calls += 1
            state.last_body = body
            text = state.text
            half = len(text) // 2
            events = [
                {"type": "message_start", "message": {"id": "msg_mock"}},
                {"type": "content_block_start", "index": 0, "content_block": {"type": "text", "text": ""}},
                {"type": "content_block_delta", "index": 0, "delta": {"type": "text_delta", "text": text[:half]}},
                {"type": "content_block_delta", "index": 0, "delta": {"type": "text_delta", "text": text[half:]}},
                {"type": "content_block_stop", "index": 0},
                {"type": "message_delta", "delta": {"stop_reason": "end_turn"}, "usage": {"output_tokens": 1}},
                {"type": "message_stop"},
            ]
            payload = "".join(f"event: {e['type']}\ndata: {json.dumps(e)}\n\n" for e in events).encode()
            self.send_response(200)
            self.send_header("content-type", "text/event-stream")
            self.send_header("content-length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)
    return Handler


def slide_text(slide):
    parts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            parts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    return "\n".join(parts)


def runs_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    yield r


def run(args, cwd=REPO):
    return subprocess.run(["node", JOB] + args, capture_output=True, text=True, cwd=cwd)


def main():
    try:
        import pptx
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.util import Emu  # noqa: F401
    except ImportError:
        print("SKIP: python-pptx not installed (pip install -r requirements.txt)")
        sys.exit(0)

    tmp = tempfile.mkdtemp(prefix="deck2pptx-gate-")
    md = os.path.join(tmp, "route-merge--draft-20260905-000000.md")
    with open(md, "w", encoding="utf-8") as f:
        f.write(DRAFT)
    spec_path = os.path.join(tmp, "deck.json")
    with open(spec_path, "w", encoding="utf-8") as f:
        json.dump(SPEC, f)
    media = os.path.join(tmp, "media")
    figs = os.path.join(tmp, "figs")
    os.makedirs(media)
    os.makedirs(figs)
    with open(os.path.join(media, "doc12_slide2_fig1.svg"), "w", encoding="utf-8") as f:
        f.write(FIG_SVG)
    with open(os.path.join(figs, "draft--fig-tc-n1.svg"), "w", encoding="utf-8") as f:
        f.write(FIG_SVG.replace("Slide 2 — Routes before the merge", "TC-N1 — Merge denied on locked route"))

    # ---- 1. the walk ---------------------------------------------------
    print("== render")
    r = run([md, "--spec", spec_path, "--media", media, "--figures", figs])
    out = md[:-3] + "--deck.pptx"
    check("renders to the sibling <stem>--deck.pptx", r.returncode == 0 and os.path.exists(out), r.stdout + r.stderr)
    check("stdout reports slides, title, proposed, dropped",
          "12 slides" in r.stdout and "Route Merge" in r.stdout and "(13 proposed, 1 dropped; design fluent, theme light)" in r.stdout, r.stdout)
    check("the invented card drops ITS slide, named on stderr with the finding",
          "dropped slide 10 (cards)" in r.stderr and "is not in the draft" in r.stderr, r.stderr)
    d = pptx.Presentation(out)
    slides = list(d.slides)
    texts = [slide_text(s) for s in slides]
    walk = ["title", "stats", "checklist", "section", "two-column", "figure", "flow", "figure", "comparison", "table", "figure", "closing"]
    check("12 slides in the spec's order", len(slides) == 12, str(len(slides)))
    check("title slide: headline, DRAFT pill, facts; stats: 'At a glance'; closing last",
          "Route Merge" in texts[0] and "DRAFT — MACHINE-GENERATED, UNREVIEWED" in texts[0] and "Claire Wang" in texts[0]
          and "At a glance" in texts[1] and "Before this draft becomes the plan" in texts[11], texts[0][:200])
    check("machine banner + raw URLs reach no slide",
          not any("machine-generated test-plan draft" in t or "mock.example" in t for t in texts), "")

    # ---- 2. the design system on the slide -----------------------------
    print("== design system")
    def bg_rgb(slide):
        bg = slide._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
        if bg is None:
            return None
        clr = bg.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        return clr.get("val") if clr is not None else None
    grounds = [bg_rgb(s) for s in slides]
    check("title / section / closing on ink (16302F), the rest on paper",
          all((g == "16302F") == (w in ("title", "section", "closing")) for g, w in zip(grounds, walk)), grounds)
    title_runs = [r for r in runs_of(slides[1]) if r.text == "At a glance"]
    check("slide title at Fluent title1 × 1.5 = 36 pt, semibold",
          title_runs and title_runs[0].font.size.pt == 36 and title_runs[0].font.bold, [(r.font.size, r.font.bold) for r in title_runs])
    step_runs = [r for r in runs_of(slides[4]) if "Run Merge Routes" in r.text]
    check("steps at body2 × 1.5 = 18 pt", step_runs and step_runs[0].font.size.pt == 18, [r.font.size for r in step_runs])
    label_runs = [r for r in runs_of(slides[4]) if r.text == "EXPECTED RESULT"]
    check("card labels at caption1Strong × 1.5 = 13.5 pt, bold",
          label_runs and label_runs[0].font.size.pt == 13.5 and label_runs[0].font.bold, [r.font.size for r in label_runs])
    fonts = {r.font.name for s in slides for r in runs_of(s) if r.font.name}
    check("Segoe UI on every run", fonts == {"Segoe UI"}, fonts)
    check("footer (plan title + page number) on paper slides only",
          all(("Test Plan — Route Merge" in texts[i]) == (walk[i] not in ("title", "section", "closing")) for i in range(12))
          and "\n5" in texts[4], "")
    W, H = 12192000, 6858000
    inside = all(sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= W + 1 and sh.top + sh.height <= H + 1
                 for s in slides for sh in s.shapes if sh.width is not None)
    check("every shape inside the 16:9 canvas", inside, "")

    # ---- 3. native objects --------------------------------------------
    print("== native objects")
    case = slides[4]
    boxes = [sh for sh in case.shapes if sh.name == "checkbox"]
    check("two-column: the steps as checkbox shapes + verbatim text",
          len(boxes) == 2 and "1. Run Merge Routes on route R1 and route R2." in texts[4] and "STEPS" in texts[4], texts[4][:300])
    cards = [sh for sh in case.shapes if sh.name in ("Expected result", "Trace")]
    check("two-column: Expected Result + Trace as filled cards on the right, TC-P1 eyebrow",
          len(cards) == 2 and all(c.left > W / 2 for c in cards) and "TC-P1 · POSITIVE" in texts[4]
          and "keeps the source measures unchanged" in texts[4], [(c.name, c.left) for c in cards])
    tiles = [sh for sh in slides[1].shapes if sh.name == "tile"]
    check("stats: four tiles with the deterministic counts (1, 1, 2, 2)",
          len(tiles) == 4 and re.search(r"\n1\nPositive cases", texts[1]) and re.search(r"\n2\nOpen \[VERIFY\] flags", texts[1]), texts[1][:300])
    check("stats: scope lede + verifier callout", "SCOPE" in texts[1] and "DRAFT VERIFIER" in texts[1] and "TC-N1 carries a" in texts[1], texts[1])
    tables = [sh for sh in slides[9].shapes if getattr(sh, "has_table", False) and sh.has_table]
    check("table: the Coverage Map as ONE native table with its cells",
          len(tables) == 1 and tables[0].table.rows[1].cells[2].text.strip() == "TC-P1"
          and "the merge must preserve measures" in tables[0].table.rows[1].cells[1].text, "")
    flow = slides[6]
    auto = [sh for sh in flow.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.name == "step"]
    prsts = [sh._element.spPr.prstGeom.get("prst") for sh in auto]
    check("flow: homePlate + chevron autoshapes, the step text INSIDE each shape",
          prsts == ["homePlate", "chevron", "chevron"] and auto[0].text_frame.text == "As User A, lock route R1."
          and auto[2].text_frame.text == "The merge is denied with a lock conflict.", prsts)
    check("flow: numbered circles + the outcome card", "EXPECTED" in texts[6] and [sh for sh in flow.shapes if sh.name == "number"], texts[6])
    comp = slides[8]
    chips = [sh for sh in comp.shapes if sh.name.startswith("chip ")]
    check("comparison: two labelled panels with pills", len(chips) == 2 and {c.name for c in chips} == {"chip TC-P1", "chip TC-N1"}
          and "denied with a lock conflict" in texts[8], [c.name for c in chips])
    section = slides[3]
    check("section: number + headline + strap on ink", "01" in texts[3] and "Positive Tests" in texts[3] and "one case" in texts[3], texts[3])

    # ---- 4. figures ----------------------------------------------------
    print("== figures")
    def groups(slide):
        return [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.GROUP]
    g5 = groups(slides[5])
    check("story figure under --media: ONE native shape group (plate dropped → 6 shapes), named by its SVG title",
          len(g5) == 1 and len(g5[0].shapes) == 6 and g5[0].name == "Slide 2 — Routes before the merge", [(g.name, len(g.shapes)) for g in g5])
    check("figure slide: title, source eyebrow, reading notes beside the group",
          "Routes before the merge" in texts[5] and "TC-P1" in texts[5] and "READING NOTES" in texts[5]
          and "Run Merge Routes on route R1 and route R2." in texts[5], texts[5][:300])
    fig_span8 = g5[0].left + g5[0].width <= 8300000  # inside the 8-column figure region when an aside is present
    check("figure with an aside stays inside the 8-column region", fig_span8, (g5[0].left, g5[0].width))
    g7 = groups(slides[7])
    check("generated figure under --figures: a native group named by the case", len(g7) == 1 and g7[0].name.startswith("TC-N1"), [g.name for g in g7])
    g10 = groups(slides[10])
    check("inline figurespec: grounded, rendered, embedded as a native group",
          len(g10) == 1 and g10[0].name == "TC-P1 — Merge preserves measures" and len(g10[0].shapes) >= 3, [(g.name, len(g.shapes)) for g in g10])

    # no --media: degrade with a note + coaching
    out2 = os.path.join(tmp, "nomedia.pptx")
    r = run([md, "--spec", spec_path, "-o", out2])
    nd = pptx.Presentation(out2)
    ntexts = [slide_text(s) for s in nd.slides]
    check("without --media / --figures the deck still renders; the figure slide carries the (not embedded) note",
          r.returncode == 0 and len(ntexts) == 12 and "Figure: Routes R1 and R2 before the merge (not embedded)" in ntexts[5]
          and not groups(list(nd.slides)[5]), r.stderr)
    check("stderr names the fix", "pass --media / --figures" in r.stderr, r.stderr)

    # ---- 5. notes, 6. amber -------------------------------------------
    print("== notes + flags")
    check("speaker notes become a native notes page",
          slides[0].has_notes_slide and slides[0].notes_slide.notes_text_frame.text == "Open with the story context; the PE owns the plan."
          and not slides[3].has_notes_slide, "")
    amber = [r.text for s in slides for r in runs_of(s)
             if r.font.color and r.font.color.type is not None and str(r.font.color.rgb) == "C2701A" and "[VERIFY:" in r.text]
    check("[VERIFY: …] spans inside copied text surface as amber runs",
          amber == ["[VERIFY: minimum network configuration]"], amber)
    check("closing provenance names the design", "design Fluent 2" in texts[11], texts[11][-200:])

    # ---- 9. designs ----------------------------------------------------
    print("== designs")
    def deck_on(design):
        outp = os.path.join(tmp, f"{design}.pptx")
        rr = run([md, "--spec", spec_path, "-o", outp, "--media", media, "--design", design])
        dd = pptx.Presentation(outp)
        ss = list(dd.slides)
        return rr, ss, [slide_text(x) for x in ss]
    rr, cs, ct = deck_on("carbon")
    cfonts = {r.font.name for s in cs for r in runs_of(s) if r.font.name}
    ctitle = [r for r in runs_of(cs[1]) if r.text == "At a glance"]
    ccard = [sh for sh in cs[4].shapes if sh.name == "Expected result"][0]
    cadj = ccard._element.spPr.prstGeom.find("{http://schemas.openxmlformats.org/drawingml/2006/main}avLst/{http://schemas.openxmlformats.org/drawingml/2006/main}gd").get("fmla")
    check("carbon: 12 slides, IBM Plex Sans on every run, heading-05 title = 36 pt regular, Gray 100 inverse",
          rr.returncode == 0 and len(cs) == 12 and cfonts == {"IBM Plex Sans"} and ctitle and ctitle[0].font.size.pt == 36
          and not ctitle[0].font.bold and bg_rgb(cs[0]) == "161616", (cfonts, [(r.font.size, r.font.bold) for r in ctitle], bg_rgb(cs[0])))
    check("carbon: square cards (roundRect adj 0), provenance names IBM Carbon, the figure group still embeds",
          cadj == "val 0" and "design IBM Carbon" in ct[11] and len(groups(cs[5])) == 1, (cadj, ct[11][-120:]))
    check("carbon: every shape inside the canvas",
          all(sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= W + 1 and sh.top + sh.height <= H + 1 for s in cs for sh in s.shapes if sh.width is not None), "")
    rr, us, ut = deck_on("uswds")
    ufonts = {r.font.name for s in us for r in runs_of(s) if r.font.name}
    utitle = [r for r in runs_of(us[1]) if r.text == "At a glance"]
    check("uswds: Public Sans on every run, size-13 title = 40.5 pt bold, primary-darker inverse, provenance names the system",
          rr.returncode == 0 and len(us) == 12 and ufonts == {"Public Sans"} and utitle and utitle[0].font.size.pt == 40.5
          and utitle[0].font.bold and bg_rgb(us[0]) == "162E51" and "design U.S. Web Design System" in ut[11],
          (ufonts, [(r.font.size, r.font.bold) for r in utitle], bg_rgb(us[0])))
    check("uswds: every shape inside the canvas",
          all(sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= W + 1 and sh.top + sh.height <= H + 1 for s in us for sh in s.shapes if sh.width is not None), "")
    rr = run([md, "--spec", spec_path, "-o", os.path.join(tmp, "bogus.pptx"), "--design", "bogus"])
    check("--design bogus refused up front, nothing written",
          rr.returncode == 2 and 'unknown design "bogus"' in rr.stderr and not os.path.exists(os.path.join(tmp, "bogus.pptx")), rr.stderr)
    rr = run([md, "--help"])
    check("--help lists the designs with licence and font, and the themes",
          rr.returncode == 0 and "carbon: IBM Carbon (Apache-2.0)" in rr.stdout and "Public Sans" in rr.stdout and "themes: light | dark" in rr.stdout, rr.stdout[-300:])

    # ---- 10. themes ----------------------------------------------------
    print("== themes")
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    def group_colours(g):
        return sorted({el.get("val") for el in g._element.iter(A + "srgbClr")})
    light_fig = group_colours(g5[0])
    outd = os.path.join(tmp, "carbon-dark.pptx")
    rr = run([md, "--spec", spec_path, "-o", outd, "--media", media, "--design", "carbon", "--theme", "dark"])
    dd = pptx.Presentation(outd)
    ds_ = list(dd.slides)
    dt = [slide_text(x) for x in ds_]
    title_dark = [r for r in runs_of(ds_[1]) if r.text == "At a glance"]
    check("carbon dark: paper slides on Gray 100, title text Gray 10, dividers on Blue 80, provenance names the theme",
          rr.returncode == 0 and bg_rgb(ds_[1]) == "161616" and str(title_dark[0].font.color.rgb) == "F4F4F4" and bg_rgb(ds_[3]) == "002D9C"
          and "design IBM Carbon (dark) · theme dark" in dt[11], (bg_rgb(ds_[1]), bg_rgb(ds_[3]), dt[11][-120:]))
    tbl = [sh for sh in ds_[9].shapes if getattr(sh, "has_table", False) and sh.has_table][0]
    head_fill = tbl._element.find(".//" + A + "tcPr/" + A + "solidFill/" + A + "srgbClr").get("val")
    check("carbon dark: the native table's header sits on the divider surface (Blue 80), not the light palette's ink",
          head_fill == "002D9C", head_fill)
    dark_fig = group_colours(groups(ds_[5])[0])
    check("carbon dark: the embedded figure is re-coloured — Blue 50 for teal, Gray 10 for ink, no Diagram Style Framework value left",
          "4589FF" in dark_fig and "F4F4F4" in dark_fig and not set(light_fig) & set(dark_fig), (light_fig, dark_fig))
    check("fluent light: the figure group keeps the palette it was drawn in (ink, teal, muted)",
          light_fig == ["16302F", "1B6E8C", "6E8285"], light_fig)
    rr = run([md, "--spec", spec_path, "-o", os.path.join(tmp, "dusk.pptx"), "--theme", "dusk"])
    check("--theme dusk refused up front", rr.returncode == 2 and 'unknown theme "dusk"' in rr.stderr and not os.path.exists(os.path.join(tmp, "dusk.pptx")), rr.stderr)

    # ---- 7. CLI --------------------------------------------------------
    print("== CLI")
    r = run([])
    check("no args: usage, nonzero", r.returncode != 0 and "usage:" in r.stderr, r.stderr[:100])
    r = run([md])
    check("neither --spec nor --generate: usage, nonzero", r.returncode != 0, "")
    r = run([md, "--spec", spec_path, "--media", os.path.join(tmp, "nope")])
    check("--media non-directory refused", r.returncode != 0 and "not a directory" in r.stderr, r.stderr)
    wrapped = os.path.join(tmp, "reply.txt")
    with open(wrapped, "w", encoding="utf-8") as f:
        f.write(DECK_REPLY)
    out3 = os.path.join(tmp, "wrapped.pptx")
    r = run([md, "--spec", wrapped, "-o", out3])
    check("--spec accepts a sentinel-wrapped reply file; -o names the output",
          r.returncode == 0 and os.path.exists(out3) and len(list(pptx.Presentation(out3).slides)) == 6, r.stderr)
    bad = os.path.join(tmp, "bad.json")
    with open(bad, "w", encoding="utf-8") as f:
        json.dump({"slides": [{"pattern": "bullets", "title": "x", "regions": {"items": ["nothing here is in the draft"]}}]}, f)
    r = run([md, "--spec", bad, "-o", os.path.join(tmp, "bad.pptx")])
    check("no surviving slide: exit 1, nothing written",
          r.returncode == 1 and "no slide survived" in r.stderr and not os.path.exists(os.path.join(tmp, "bad.pptx")), r.stderr)

    # ---- 8. --generate against a mock ---------------------------------
    print("== generate")
    state = MockState()
    server = ThreadingHTTPServer(("127.0.0.1", 0), make_handler(state))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{server.server_port}"
    cfg = os.path.join(tmp, "config.json")
    with open(cfg, "w", encoding="utf-8") as f:
        json.dump({"llm": {"provider": "anthropic", "apiKey": "mock-key", "baseUrl": base, "maxRetries": 0},
                   "testplangen": {"deckMaxTokens": 12345, "deckDesign": "carbon"}}, f)
    out4 = os.path.join(tmp, "generated.pptx")
    r = run([md, "--generate", "--config", cfg, "-o", out4, "--media", media])
    prompt = (state.last_body.get("messages") or [{}])[0].get("content", "")
    check("one streamed model call; the repo prompt verbatim with the inputs substituted",
          r.returncode == 0 and state.calls == 1 and state.last_body.get("stream") is True
          and "DECK SPECIFICATION VOCABULARY" in prompt and "<<<DRAFT BEGIN>>>" in prompt
          and "### TC-P1 — Merge preserves measures" in prompt and "The plan title: Test Plan — Route Merge" in prompt
          and not re.search(r"\{(PlanTitle|Draft|Figures)\}", prompt), r.stderr[-400:] + prompt[:200])
    check("the Figures input lists the cited story + generated figures by file name",
          "- doc12_slide2_fig1.svg — story figure for TC-P1: Routes R1 and R2 before the merge" in prompt
          and "- draft--fig-tc-n1.svg — generated figure for TC-N1" in prompt, prompt[prompt.find("The figures"):][:300])
    check("max_tokens = testplangen.deckMaxTokens", state.last_body.get("max_tokens") == 12345, state.last_body.get("max_tokens"))
    spec_out = out4[:-5] + ".deck.json"
    gen_deck = pptx.Presentation(out4)
    check("the spec lands beside the deck as <out>.deck.json; the deck has the reply's 6 slides",
          os.path.exists(spec_out) and json.load(open(spec_out))["slides"][0]["pattern"] == "title"
          and len(list(gen_deck.slides)) == 6, r.stderr)
    check("--generate takes the design from testplangen.deckDesign (carbon → IBM Plex Sans)",
          {r.font.name for s in gen_deck.slides for r in runs_of(s) if r.font.name} == {"IBM Plex Sans"}, "")
    r = run([md, "--spec", spec_out, "-o", os.path.join(tmp, "rerender.pptx"), "--media", media])
    check("the written spec re-renders with --spec", r.returncode == 0 and "6 slides" in r.stdout, r.stdout + r.stderr)
    state.text = "no sentinels here"
    r = run([md, "--generate", "--config", cfg, "-o", os.path.join(tmp, "nosent.pptx")])
    check("sentinel-less reply: exit nonzero, nothing written",
          r.returncode != 0 and "DECK BEGIN/END" in r.stderr and not os.path.exists(os.path.join(tmp, "nosent.pptx")), r.stderr)
    server.shutdown()

    print(f"\n{len(PASS)}/{len(PASS) + len(FAIL)} checks passed")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
